#!/usr/bin/env python3
"""Generate synthetic fixture decks with known accessibility issues.

All content is synthetic and generated here, so the decks may be redistributed under
the repository's MIT license. Requires python-pptx (and Pillow, which it depends on).

    python3 fixtures/generate.py

Expected inventories for each deck live in fixtures/expected/*.json and are asserted
by tests/fixtures.test.ts. Regenerating must not change those expectations.
"""
from __future__ import annotations

import copy
import io
import os
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "decks"
NSMAP = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
FIXED_TIME = (2026, 9, 1, 0, 0, 0)


def png(color: tuple[int, int, int], label: str, size=(640, 400), alpha=False) -> io.BytesIO:
    mode = "RGBA" if alpha else "RGB"
    img = Image.new(mode, size, color + ((0,) if alpha else ()))
    d = ImageDraw.Draw(img)
    if alpha:
        d.ellipse((80, 40, size[0] - 80, size[1] - 40), fill=color + (255,))
    else:
        d.rectangle((20, 20, size[0] - 20, size[1] - 20), outline=(255, 255, 255), width=8)
    d.text((40, size[1] // 2), label, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def jpeg(label: str) -> io.BytesIO:
    img = Image.new("RGB", (800, 600))
    d = ImageDraw.Draw(img)
    for y in range(600):
        d.line((0, y, 800, y), fill=(30 + y // 5, 90, 160 - y // 6))
    d.text((40, 40), label, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    return buf


def set_descr(shape, descr: str | None = None, title: str | None = None) -> None:
    cNvPr = shape._element.xpath("./*[1]/p:cNvPr")[0]
    for k in ("descr", "title"):
        if k in cNvPr.attrib:
            del cNvPr.attrib[k]
    if descr is not None:
        cNvPr.set("descr", descr)
    if title is not None:
        cNvPr.set("title", title)


def clear_descr(shape) -> None:
    set_descr(shape, None, None)


def add_link_text(slide, text_before: str, link_text: str, url: str, top) -> None:
    tb = slide.shapes.add_textbox(Inches(0.5), top, Inches(8), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = text_before
    r2 = p.add_run()
    r2.text = link_text
    r2.hyperlink.address = url


def build_basic() -> Presentation:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_only = prs.slide_layouts[5]
    title_content = prs.slide_layouts[1]

    # 1. Title slide (ctrTitle), no issues besides reading order review.
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Introduction to Accessible Slides"
    s.placeholders[1].text = "Synthetic fixture deck"

    # 2. Picture missing alt text + picture with good alt text; title placeholder.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Enrollment trends"
    shared = png((40, 90, 160), "shared media")
    pic = s.shapes.add_picture(shared, Inches(0.5), Inches(1.5), Inches(4), Inches(2.5))
    clear_descr(pic)
    good = s.shapes.add_picture(jpeg("lab photo"), Inches(5), Inches(1.5), Inches(4), Inches(3))
    set_descr(good, "Students collaborate at computers in a teaching lab.")

    # 3. Same media file reused (different context), generic alt text, filename alt.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Campus resources"
    shared.seek(0)
    reuse = s.shapes.add_picture(shared, Inches(0.5), Inches(1.5), Inches(3), Inches(2))
    clear_descr(reuse)
    generic = s.shapes.add_picture(png((150, 60, 60), "generic"), Inches(4), Inches(1.5), Inches(2.5), Inches(2))
    set_descr(generic, "image1.png")
    autogen = s.shapes.add_picture(png((60, 140, 60), "auto"), Inches(7), Inches(1.5), Inches(2.5), Inches(2))
    set_descr(autogen, "A picture containing text\n\nDescription automatically generated")

    # 4. No title at all (blank layout) with text and a title-only legacy alt field.
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
    tb.text_frame.text = "This looks like a title but is only a text box"
    legacy = s.shapes.add_picture(png((90, 90, 90), "legacy title"), Inches(1), Inches(2), Inches(3), Inches(2))
    set_descr(legacy, None, "Legacy title only")

    # 5 & 6. Duplicate titles.
    for _ in range(2):
        s = prs.slides.add_slide(title_content)
        s.shapes.title.text = "Review Questions"
        s.placeholders[1].text = "What did we learn?"

    # 7. Links: generic text and naked URL.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Further reading"
    add_link_text(s, "For the syllabus ", "click here", "https://example.edu/syllabus", Inches(1.6))
    add_link_text(s, "Library: ", "https://library.example.edu/guides", "https://library.example.edu/guides", Inches(2.4))
    add_link_text(s, "See the ", "accessibility guidelines", "https://example.edu/a11y", Inches(3.2))

    # 8. Table without header row, and a table with merged cells.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Grading breakdown"
    gf = s.shapes.add_table(3, 2, Inches(0.5), Inches(1.5), Inches(4), Inches(1.5))
    tbl = gf.table
    for r, row in enumerate([["Item", "Weight"], ["Exams", "60%"], ["Labs", "40%"]]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    tbl.first_row = False
    gf2 = s.shapes.add_table(3, 3, Inches(5), Inches(1.5), Inches(4.5), Inches(1.5))
    t2 = gf2.table
    for r in range(3):
        for c in range(3):
            t2.cell(r, c).text = f"R{r}C{c}"
    t2.cell(0, 0).merge(t2.cell(0, 1))

    # 9. Chart (graphicFrame) without alt text.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Survey results"
    data = CategoryChartData()
    data.categories = ["Agree", "Neutral", "Disagree"]
    data.add_series("Responses", (42, 17, 9))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(7), Inches(4.5), data)

    # 10. Group containing a picture; cropped picture; low contrast text.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Lab safety"
    grp = s.shapes.add_group_shape()
    gp = grp.shapes.add_picture(png((200, 120, 20), "grouped"), Inches(0.5), Inches(1.5), Inches(2), Inches(2))
    clear_descr(gp)
    grp.shapes.add_textbox(Inches(0.5), Inches(3.6), Inches(2), Inches(0.5)).text_frame.text = "Goggles"
    cropped = s.shapes.add_picture(png((20, 120, 200), "cropped"), Inches(3.5), Inches(1.5), Inches(3), Inches(2))
    cropped.crop_left = 0.1
    cropped.crop_right = 0.2
    clear_descr(cropped)
    low = s.shapes.add_textbox(Inches(7), Inches(1.5), Inches(2.5), Inches(1))
    low.fill.solid()
    low.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run = low.text_frame.paragraphs[0].add_run()
    run.text = "Faint warning text"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    ok = s.shapes.add_textbox(Inches(7), Inches(3), Inches(2.5), Inches(1))
    ok.fill.solid()
    ok.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run = ok.text_frame.paragraphs[0].add_run()
    run.text = "Readable text"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # 11. Hidden slide with transparent PNG and notes.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Instructor notes"
    tp = s.shapes.add_picture(png((10, 150, 90), "transparent", alpha=True), Inches(1), Inches(1.5), Inches(3), Inches(2))
    clear_descr(tp)
    s.notes_slide.notes_text_frame.text = "Remember to pause for questions."
    s._element.set("show", "0")

    # 12. Unicode / RTL title and picture with quotes/ampersand-worthy context.
    s = prs.slides.add_slide(title_only)
    s.shapes.title.text = "Café & “Résumé” — مرحبا"
    up = s.shapes.add_picture(png((120, 40, 140), "unicode"), Inches(1), Inches(1.5), Inches(3), Inches(2))
    clear_descr(up)
    return prs


def build_reordered() -> Presentation:
    prs = Presentation()
    layout = prs.slide_layouts[5]
    for name in ["First", "Second", "Third"]:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = name
    # Reverse presentation order without renaming slide parts.
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    for el in ids:
        sldIdLst.remove(el)
    for el in reversed(ids):
        sldIdLst.append(el)
    return prs


def build_sample() -> Presentation:
    """Public demo deck with deliberate issues (spec §23)."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Sample Deck: Photosynthesis"
    s.placeholders[1].text = "A synthetic presentation with deliberate accessibility issues"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "How plants make food"
    pic = s.shapes.add_picture(png((30, 130, 60), "leaf diagram"), Inches(0.5), Inches(1.5), Inches(4.5), Inches(3))
    clear_descr(pic)
    tb = s.shapes.add_textbox(Inches(5.3), Inches(1.5), Inches(4.2), Inches(3))
    tb.text_frame.text = "Light energy + water + carbon dioxide → glucose + oxygen"
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
    t.text_frame.text = "Where it happens"
    p2 = s.shapes.add_picture(jpeg("chloroplast"), Inches(1), Inches(1.6), Inches(5), Inches(3.75))
    set_descr(p2, "chloroplast.jpg")
    add_link_text(s, "Learn more: ", "click here", "https://example.org/photosynthesis", Inches(5.6))
    return prs


def normalize_zip_times(path: Path) -> None:
    """Rewrite the archive with fixed timestamps so fixtures are reproducible."""
    import zipfile

    with zipfile.ZipFile(path) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(zi, data)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, builder in [("basic", build_basic), ("reordered", build_reordered), ("sample-deck", build_sample)]:
        path = OUT / f"{name}.pptx"
        builder().save(path)
        normalize_zip_times(path)
        print(f"wrote {path.relative_to(Path.cwd()) if path.is_relative_to(Path.cwd()) else path}")


if __name__ == "__main__":
    main()
